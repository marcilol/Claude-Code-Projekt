# -*- coding: utf-8 -*-
"""Generate Portfolio X-Ray Project Plan as Word document."""

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import datetime


def set_cell_shading(cell, color):
    """Set background color of a table cell."""
    shading = cell._element.get_or_add_tcPr()
    shd = shading.makeelement(qn('w:shd'), {
        qn('w:fill'): color,
        qn('w:val'): 'clear',
    })
    shading.append(shd)


def add_table(doc, headers, rows, col_widths=None):
    """Add a formatted table to the document."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            p.style = doc.styles['Normal']
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)
        set_cell_shading(cell, '2F5496')
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.color.rgb = RGBColor(255, 255, 255)

    # Data rows
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
            if r_idx % 2 == 1:
                set_cell_shading(cell, 'D6E4F0')

    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Cm(w)

    return table


def build_document():
    doc = Document()

    # -- Title Page --
    doc.add_paragraph()
    doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('Portfolio X-Ray')
    run.font.size = Pt(36)
    run.bold = True
    run.font.color.rgb = RGBColor(47, 84, 150)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run('Factor-Based Portfolio Risk Analysis Tool')
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(89, 89, 89)

    doc.add_paragraph()
    subtitle2 = doc.add_paragraph()
    subtitle2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle2.add_run('Project Plan & Technical Overview')
    run.font.size = Pt(16)

    doc.add_paragraph()
    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = date_p.add_run(f'{datetime.date.today().strftime("%B %d, %Y")}')
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(89, 89, 89)

    doc.add_page_break()

    # -- Table of Contents placeholder --
    doc.add_heading('Table of Contents', level=1)
    toc_items = [
        '1. Current Model Overview',
        '2. Verification & Cross-Checking Plan',
        '3. Front-End Architecture',
        '4. Database & Data Pipeline',
        '5. Data Provider Recommendation',
        '6. International Market Expansion',
        '7. Industry Classification Upgrade',
        '8. Additional Considerations',
        '9. Development Roadmap',
        '10. GitHub Collaboration Guide',
    ]
    for item in toc_items:
        p = doc.add_paragraph(item)
        p.paragraph_format.space_after = Pt(2)

    doc.add_page_break()

    # ================================================================
    # SECTION 1: Current Model Overview
    # ================================================================
    doc.add_heading('1. Current Model Overview', level=1)

    doc.add_heading('1.1 What the Tool Does', level=2)
    doc.add_paragraph(
        'Portfolio X-Ray is a Python-based factor risk analysis tool that implements '
        'the MSCI Barra CNE5 cross-sectional regression methodology. It decomposes '
        'portfolio risk into systematic (factor-driven) and idiosyncratic (stock-specific) '
        'components, based on Giuseppe Paleologo\'s "Advanced Portfolio Management."'
    )

    doc.add_heading('1.2 Current Capabilities', level=2)
    features = [
        ('Data Pipeline', 'Fetches 2 years of daily price/fundamental data for ~2,545 Russell 3000 stocks from Yahoo Finance. Computes 10 CNE5-style style factors + 11 GICS sector factors + 1 Country factor = 22 total factors.'),
        ('Factor Model', 'Daily cross-sectional WLS regression (sqrt market cap weights) estimating factor returns. Newey-West autocorrelation adjustment, eigenfactor risk adjustment (Monte Carlo), and volatility regime adjustment. Average R\u00b2 \u2248 28% (WLS-weighted).'),
        ('Risk Decomposition', 'Breaks portfolio volatility into factor risk (market, industry, style) vs. idiosyncratic risk. Identifies which factors drive portfolio risk.'),
        ('Factor-Neutral Optimization', 'Quadratic programming to minimize factor exposure while staying close to original weights. Long-only constraint.'),
        ('Alpha Sizing', 'Converts expected returns into position sizes using four methods: Proportional, Risk Parity, Mean-Variance, and Shrunk MV. Includes backtesting framework.'),
        ('Factor Risk Management', 'MCFR (Marginal Contribution to Factor Risk) analysis, risk limit checking (idiosyncratic %, concentration, HHI), and trade suggestions.'),
    ]
    add_table(doc, ['Capability', 'Description'], features, col_widths=[4, 13])

    doc.add_heading('1.3 Factor Model Details', level=2)
    doc.add_paragraph(
        'The model implements 10 CNE5-style style factors, each computed from '
        'sub-descriptors following MSCI Barra methodology:'
    )
    factors = [
        ('Size', 'LNCAP', 'log(market cap)'),
        ('Beta', 'BETA', 'EWM regression vs market (halflife=63)'),
        ('Momentum', 'RSTR', 'EWM of lagged excess log returns (skip 21d, hl=126)'),
        ('Residual Vol', '0.74*DASTD + 0.16*CMRA + 0.10*HSIGMA', 'Orthogonalized vs beta & size'),
        ('NL Size', 'Cube of z(size)', 'Orthogonalized vs size'),
        ('Book-to-Price', 'BTOP', 'Book value / market cap'),
        ('Liquidity', '0.35*STOM + 0.35*STOQ + 0.30*STOA', 'Orthogonalized vs size'),
        ('Earnings Yield', '0.656*CETOP + 0.344*ETOP', 'Partial: missing EPFWD (68% weight)'),
        ('Growth', '0.338*EGRO + 0.662*SGRO', 'Partial: missing analyst forecasts (29% weight)'),
        ('Leverage', '0.38*MLEV + 0.35*DTOA + 0.27*BLEV', 'Full implementation'),
    ]
    add_table(doc, ['Factor', 'Descriptor(s)', 'Notes'], factors, col_widths=[3, 6.5, 7.5])

    doc.add_heading('1.4 Model Adjustments', level=2)
    adjustments = [
        ('WLS Regression', 'sqrt(market cap) weights for cross-sectional regression'),
        ('Industry Neutrality', 'Cap-weighted constraint to handle Country/industry multicollinearity'),
        ('Newey-West', 'Autocorrelation adjustment (MA order q=2, halflife=252)'),
        ('Eigenfactor Risk Adj.', 'Monte Carlo simulation (100 sims, scale=1.4) corrects eigenvalue bias'),
        ('Volatility Regime Adj.', 'Adaptive scaling (halflife=42) based on realized vs predicted volatility'),
        ('Factor Exposure Lagging', 'All exposures measured at t-1 close to avoid look-ahead bias'),
        ('Log Returns', 'Dependent variable uses log returns for theoretical consistency'),
        ('Excess Returns for Beta', 'Beta estimated from excess returns per CNE5 specification'),
    ]
    add_table(doc, ['Adjustment', 'Description'], adjustments, col_widths=[4.5, 12.5])

    doc.add_heading('1.5 Current Limitations', level=2)
    limitations = [
        'Missing analyst consensus forecasts (EPFWD for Earnings Yield, EGRLF/EGRSF for Growth)',
        'Only 11 GICS sectors as industry factors (CNE5 uses 32 industry groups)',
        'Yahoo Finance data quality: gaps, stale fundamentals, unreliable for smaller stocks',
        'US equities only (Russell 3000)',
        'No specific risk model with Bayesian shrinkage (CNE5 feature)',
        'No short-term reversal factor',
        '1-year estimation window (vs 2-5 years in production models)',
        'CLI-only interface (no GUI)',
    ]
    for item in limitations:
        doc.add_paragraph(item, style='List Bullet')

    doc.add_page_break()

    # ================================================================
    # SECTION 2: Verification & Cross-Checking
    # ================================================================
    doc.add_heading('2. Verification & Cross-Checking Plan', level=1)

    doc.add_paragraph(
        'Since investment decisions depend on these results, the model must be '
        'thoroughly validated. We propose a multi-layered verification approach.'
    )

    doc.add_heading('2.1 Excel Verification Model', level=2)
    doc.add_paragraph(
        'Build a simplified Excel workbook that replicates the core regression for '
        'a small universe (~20-30 stocks, 1 month of data). This allows manual '
        'inspection of every intermediate calculation.'
    )
    excel_sheets = [
        ('Raw Data', 'Daily prices, market cap, fundamentals for 20-30 stocks'),
        ('Factor Exposures', 'Compute each descriptor step-by-step with visible formulas'),
        ('Z-Scoring', 'Cap-weighted mean, EW std normalization (show each step)'),
        ('Cross-Sectional Reg.', 'WLS regression for one date using Excel LINEST or matrix formulas'),
        ('Factor Returns', 'Compare Excel factor returns vs Python output'),
        ('Risk Decomposition', 'Portfolio factor variance, idio variance, total vol'),
        ('Reconciliation', 'Side-by-side comparison of Excel vs Python for every output'),
    ]
    add_table(doc, ['Sheet', 'Contents'], excel_sheets, col_widths=[4.5, 12.5])

    doc.add_heading('2.2 Automated Test Suite', level=2)
    tests = [
        ('Unit Tests', 'Test each function in isolation: z-scoring, beta computation, composites, orthogonalization. Use known inputs with hand-calculated expected outputs.'),
        ('Regression Tests', 'Save current model outputs as baselines. After any code change, compare new outputs to baselines. Flag deviations > threshold.'),
        ('Sanity Checks', 'Factor exposures: CW mean \u2248 0, EW std \u2248 1 per date. Factor returns: reasonable magnitudes (< 5% daily). R\u00b2: 20-35% range. Covariance matrix: positive definite.'),
        ('Cross-Validation', 'Split time series in half. Estimate model on first half, predict risk on second half. Check bias statistics (should be near 1.0).'),
    ]
    add_table(doc, ['Test Type', 'Description'], tests, col_widths=[3.5, 13.5])

    doc.add_heading('2.3 External Benchmarking', level=2)
    doc.add_paragraph(
        'Compare model outputs against known benchmarks where possible:'
    )
    benchmarks = [
        'Factor return signs and magnitudes vs. published factor research (e.g., Fama-French)',
        'Risk decomposition of S&P 500 ETF (SPY) should show ~95%+ factor risk, minimal idiosyncratic',
        'Risk decomposition of an equal-weight random 100-stock portfolio should show high idiosyncratic risk',
        'Beta factor return should correlate highly with market excess return',
        'Size factor return should correlate with Russell 2000 minus Russell 1000 spread',
    ]
    for item in benchmarks:
        doc.add_paragraph(item, style='List Bullet')

    doc.add_page_break()

    # ================================================================
    # SECTION 3: Front-End Architecture
    # ================================================================
    doc.add_heading('3. Front-End Architecture', level=1)

    doc.add_heading('3.1 Recommended Approach: Streamlit \u2192 Dash', level=2)
    doc.add_paragraph(
        'Start with Streamlit for rapid prototyping (1-3 weeks to MVP), '
        'then evaluate whether to migrate to Dash for more polished interactivity.'
    )

    comparison = [
        ('Streamlit', 'Python-only, fastest to build, free cloud hosting', 'Re-runs entire script on interaction, limited layout control, no background tasks', 'MVP / demo'),
        ('Dash (Plotly)', 'Superior tables, callback architecture, multi-page apps, Flask-based', 'Steeper learning curve, more boilerplate', 'Production dashboard'),
        ('React + FastAPI', 'Maximum flexibility, clean API separation, best for teams', '2x technology surface (Python + JS), slowest to build', 'If scaling to external users'),
    ]
    add_table(doc, ['Framework', 'Pros', 'Cons', 'Best For'], comparison, col_widths=[3, 5.5, 5, 3.5])

    doc.add_heading('3.2 Streamlit MVP Page Structure', level=2)
    pages = [
        ('Portfolio Upload', 'Upload CSV, preview holdings, validate tickers against universe'),
        ('Risk Analysis', 'Factor vs idiosyncratic risk breakdown, factor exposure chart, position-level table'),
        ('Optimization', 'Factor-neutral optimization, before/after comparison, download optimized weights'),
        ('Alpha Sizing', 'Input expected returns, select sizing method, backtest results, position sizing output'),
        ('Risk Management', 'MCFR analysis, risk limits dashboard, suggested trades'),
        ('Model Status', 'Data freshness, R\u00b2 over time, factor return chart, coverage stats'),
    ]
    add_table(doc, ['Page', 'Features'], pages, col_widths=[4, 13])

    doc.add_page_break()

    # ================================================================
    # SECTION 4: Database & Data Pipeline
    # ================================================================
    doc.add_heading('4. Database & Data Pipeline', level=1)

    doc.add_heading('4.1 Database: DuckDB', level=2)
    doc.add_paragraph(
        'DuckDB is an embedded columnar database optimized for analytical queries. '
        'It requires zero infrastructure (no server), integrates natively with pandas, '
        'and will shrink the current ~500 MB of CSVs to ~100 MB while dramatically '
        'improving query speed.'
    )
    db_comparison = [
        ('DuckDB', 'Columnar, zero-infra, pandas-native, fast analytics', 'Recommended'),
        ('PostgreSQL', 'Production-grade RDBMS, great for concurrent writes', 'Overkill \u2014 no concurrent write needs'),
        ('SQLite', 'Simple, embedded, widely supported', 'Row-oriented: slow for analytical queries'),
        ('Parquet files', 'Columnar, no database needed, DuckDB reads directly', 'Good complement to DuckDB'),
    ]
    add_table(doc, ['Option', 'Description', 'Verdict'], db_comparison, col_widths=[3, 10, 4])

    doc.add_heading('4.2 Data Update Pipeline', level=2)
    doc.add_paragraph('Daily update flow (runs after US market close, ~6:30 PM ET):')
    pipeline_steps = [
        ('1', 'Fetch new daily prices + any updated fundamentals from data API'),
        ('2', 'Recompute factor exposures for latest date (append to history)'),
        ('3', 'Re-run cross-sectional regression for latest date (append factor returns)'),
        ('4', 'Update factor covariance matrix (expanding window Newey-West)'),
        ('5', 'Write results to DuckDB'),
        ('6', 'Log success/failure, send notification if error'),
    ]
    add_table(doc, ['Step', 'Action'], pipeline_steps, col_widths=[1.5, 15.5])

    doc.add_paragraph()
    doc.add_paragraph(
        'Scheduling: Use Windows Task Scheduler initially, migrate to APScheduler '
        '(Python library) when the web app is deployed. For cloud deployment, '
        'use a cron job or the hosting platform\'s scheduler.'
    )

    doc.add_page_break()

    # ================================================================
    # SECTION 5: Data Provider Recommendation
    # ================================================================
    doc.add_heading('5. Data Provider Recommendation', level=1)

    doc.add_paragraph(
        'The current model uses Yahoo Finance (free, via yfinance). While adequate for '
        'prototyping, it has significant limitations: unreliable API, no analyst consensus '
        'forecasts, no GICS sub-industry codes, and inconsistent fundamental data for '
        'smaller companies.'
    )

    doc.add_heading('5.1 Provider Comparison', level=2)
    providers = [
        ('EODHD', '$60-80/mo', 'Yes', 'Yes', '70+ exchanges', 'Best value \u2014 recommended'),
        ('FMP', '$149/mo', 'Yes (Ultimate)', 'No (SIC only)', 'Global', 'Good API, expensive for estimates'),
        ('Alpha Vantage', '$50-100/mo', 'Partial', 'No', '20+ exchanges', 'No bulk consensus estimates'),
        ('Polygon/Massive', '$29-199/mo', 'No', 'No', 'US-first', 'Pricing focus, wrong fit'),
        ('Tiingo', '$10-30/mo', 'No', 'No', 'US-focused', 'Good prices, missing estimates'),
        ('SimFin', '$15-71/mo', 'No', 'No', 'US-only', 'Good fundamentals, missing key data'),
        ('Intrinio', '$150+/mo', 'Yes (Zacks)', 'Unconfirmed', 'US/Canada', 'Institutional pricing'),
        ('IEX Cloud', 'N/A', 'N/A', 'N/A', 'N/A', 'Shut down Aug 2024'),
    ]
    add_table(doc,
        ['Provider', 'Cost/mo', 'Analyst Estimates', 'GICS Sub-Industry', 'Intl Coverage', 'Verdict'],
        providers, col_widths=[2.5, 2.2, 2.8, 2.8, 2.8, 4])

    doc.add_heading('5.2 Recommendation: EODHD', level=2)
    doc.add_paragraph(
        'EODHD (EOD Historical Data) at $79.99/month (annual billing) is the best single-provider '
        'solution. It uniquely offers both analyst consensus estimates AND GICS sub-industry '
        'codes \u2014 the two biggest gaps in the current model \u2014 along with global '
        'coverage across 70+ exchanges.'
    )
    doc.add_paragraph('What EODHD provides that we currently lack:')
    eodhd_benefits = [
        'Forward EPS estimates (avg, low, high, # analysts) \u2192 enables EPFWD descriptor (68% of Earnings Yield)',
        'Long-term growth forecasts \u2192 enables EGRLF descriptor (18% of Growth)',
        'GICS sub-industry classification \u2192 enables 30+ industry factors instead of 11',
        'Global equity coverage \u2192 enables international expansion',
        '100,000 API calls/day \u2192 sufficient for ~3,000 stocks daily',
    ]
    for item in eodhd_benefits:
        doc.add_paragraph(item, style='List Bullet')

    doc.add_paragraph()
    doc.add_paragraph(
        'Hybrid approach (cheapest): Keep yfinance for daily prices (free, already working). '
        'Add EODHD Fundamentals plan ($59.99/mo) for analyst estimates + GICS codes only.'
    )

    doc.add_page_break()

    # ================================================================
    # SECTION 6: International Market Expansion
    # ================================================================
    doc.add_heading('6. International Market Expansion', level=1)

    doc.add_paragraph(
        'Expanding beyond US equities requires addressing several challenges:'
    )

    challenges = [
        ('Data Availability', 'Need a data provider with global coverage (EODHD covers 70+ exchanges). Fundamental data quality varies by market \u2014 emerging markets have sparser coverage.'),
        ('Currency Risk', 'Returns must be converted to a common currency (USD) or local currency. Factor model should include a currency factor or separate models per region.'),
        ('Trading Calendars', 'Different markets have different holidays and trading hours. The common-date intersection approach in fetch_data.py won\'t work across regions.'),
        ('Industry Classification', 'GICS codes are globally consistent, which helps. But industry composition varies: e.g., Materials is 2% of S&P 500 but 15% of Australian market.'),
        ('Regulatory Differences', 'Short-selling restrictions, ownership limits, and reporting standards vary. Affects leverage and earnings descriptors.'),
        ('Factor Behavior', 'Factor premia differ across regions. Momentum is stronger in Europe, value is stronger in emerging markets. May need region-specific model calibration.'),
    ]
    add_table(doc, ['Challenge', 'Details'], challenges, col_widths=[3.5, 13.5])

    doc.add_heading('6.1 Recommended Expansion Path', level=2)
    expansion = [
        ('Phase 1', 'US Russell 3000 (current)', 'Validate and polish'),
        ('Phase 2', 'Add European developed (STOXX 600)', 'Similar data quality to US'),
        ('Phase 3', 'Add Asia-Pacific developed (ASX 200, Nikkei 225, etc.)', 'Different market structure'),
        ('Phase 4', 'Emerging markets (selectively)', 'Data quality challenges'),
    ]
    add_table(doc, ['Phase', 'Market', 'Notes'], expansion, col_widths=[2, 7, 8])

    doc.add_heading('6.2 Architecture Decision: Regional vs Global Model', level=2)
    doc.add_paragraph(
        'Two approaches exist for multi-market factor models:'
    )
    models = [
        ('Regional Models', 'Separate model per region (US, Europe, Asia). Each has its own factor returns and covariance. Simpler, allows region-specific calibration.', 'Recommended for Phase 2-3'),
        ('Global Model', 'Single model with country/region factors. More complex but captures cross-market correlations. MSCI GEM3 approach.', 'Consider for Phase 4+'),
    ]
    add_table(doc, ['Approach', 'Description', 'Recommendation'], models, col_widths=[3, 11, 3])

    doc.add_page_break()

    # ================================================================
    # SECTION 7: Industry Classification Upgrade
    # ================================================================
    doc.add_heading('7. Industry Classification Upgrade', level=1)

    doc.add_paragraph(
        'The current model uses 11 GICS sectors (top-level). MSCI Barra CNE5 uses 32 '
        'industry groups. Finer industry classification captures within-sector return '
        'dispersion (e.g., Software vs Semiconductors within Information Technology).'
    )

    doc.add_heading('7.1 GICS Hierarchy', level=2)
    gics = [
        ('Sector', '2-digit', '11', 'Current implementation'),
        ('Industry Group', '4-digit', '25', 'Good intermediate step'),
        ('Industry', '6-digit', '74', 'USE4 (US model) level'),
        ('Sub-Industry', '8-digit', '163', 'Maximum granularity'),
    ]
    add_table(doc, ['Level', 'Code', 'Count', 'Notes'], gics, col_widths=[3.5, 2.5, 2, 9])

    doc.add_heading('7.2 Where to Get GICS Codes', level=2)
    gics_sources = [
        ('EODHD', 'GicSector, GicGroup, GicIndustry, GicSubIndustry fields in fundamentals API', '$60-80/mo (included)'),
        ('Yahoo Finance', 'Provides sector and industry strings (not codes), can be mapped to GICS', 'Free but imprecise'),
        ('S&P/MSCI Direct', 'Official GICS classification database', 'Expensive (institutional)'),
        ('Wikipedia/Manual', 'GICS code tables are publicly documented; manual mapping possible', 'Free but labor-intensive'),
    ]
    add_table(doc, ['Source', 'Details', 'Cost'], gics_sources, col_widths=[3.5, 10, 3.5])

    doc.add_heading('7.3 Implementation Plan', level=2)
    doc.add_paragraph(
        'Target: move from 11 sectors to ~25-30 industry groups. This requires:'
    )
    steps = [
        'Obtain GICS industry group codes for all Russell 3000 stocks (via EODHD or mapping)',
        'Update russell_constituents.csv to include GICS industry group column',
        'Modify fetch_data.py to use industry group for dummy variables instead of sector',
        'Update industry neutrality constraint in run_factor_model.py (handles more industries)',
        'Re-estimate model and verify R\u00b2 improvement (expected +3-5%)',
    ]
    for i, item in enumerate(steps, 1):
        doc.add_paragraph(f'{i}. {item}')

    doc.add_page_break()

    # ================================================================
    # SECTION 8: Additional Considerations
    # ================================================================
    doc.add_heading('8. Additional Considerations', level=1)

    doc.add_paragraph(
        'Items not yet discussed but important for a production-quality tool:'
    )

    additional = [
        ('Specific Risk Model', 'CNE5 includes Bayesian shrinkage for idiosyncratic risk forecasts (by market-cap decile). Our model currently uses raw residual volatility. Adding shrinkage would improve risk forecasts for small-cap stocks.'),
        ('Bias Statistics / Backtesting', 'CNE5 Appendix C defines bias statistics to measure forecast accuracy. Implementing rolling bias stats would let us quantify how well the model predicts risk and track improvements over time.'),
        ('Short-Term Reversal Factor', 'A 1-week to 1-month reversal factor is a strong cross-sectional predictor absent from the current model. Easy to implement, likely +2-3% R\u00b2.'),
        ('Transaction Cost Model', 'For optimization and sizing to be practically useful, we need a transaction cost model (spread + market impact). Paleologo covers this in Chapter 5.'),
        ('Portfolio Rebalancing', 'Currently each analysis is a point-in-time snapshot. A rebalancing framework would track portfolios over time, generate periodic rebalance trades, and measure turnover.'),
        ('User Authentication', 'If the web app is public, need login/auth to protect portfolio data. Streamlit has streamlit-authenticator; Dash can use Flask-Login.'),
        ('Data Privacy', 'Portfolio holdings are sensitive. Ensure data is not logged, stored temporarily only during session, and encrypted at rest if persisted.'),
        ('Performance Optimization', 'Loading 500 MB of CSVs is slow. Migration to DuckDB/Parquet will help. For the web app, pre-compute and cache common queries.'),
        ('Error Handling & Monitoring', 'Production data pipelines need: retry logic for API failures, data quality checks (e.g., detect if a stock\'s price jumped 1000%), alerting on pipeline failures.'),
        ('Documentation', 'User-facing documentation explaining what each output means, how to interpret factor exposures, and what actions to take based on the analysis.'),
        ('Licensing', 'If distributing the tool, consider open-source license (MIT, Apache 2.0) vs proprietary. Factor model methodology is not patentable but brand names (Barra) are trademarked.'),
    ]
    add_table(doc, ['Topic', 'Details'], additional, col_widths=[4, 13])

    doc.add_page_break()

    # ================================================================
    # SECTION 9: Development Roadmap
    # ================================================================
    doc.add_heading('9. Development Roadmap', level=1)

    doc.add_heading('9.1 Phase 1: Foundation (Weeks 1-4)', level=2)
    phase1 = [
        ('Excel verification model', 'Build 20-30 stock verification workbook', 'High'),
        ('Unit test suite', 'Test core functions with known inputs/outputs', 'High'),
        ('DuckDB migration', 'Convert CSVs to DuckDB, update scripts to query DB', 'Medium'),
        ('Streamlit MVP', 'Basic upload + risk analysis page', 'High'),
        ('GitHub setup', 'Repository, branch strategy, CI/CD, issue templates', 'High'),
    ]
    add_table(doc, ['Task', 'Description', 'Priority'], phase1, col_widths=[4, 10, 3])

    doc.add_heading('9.2 Phase 2: Data Quality (Weeks 5-8)', level=2)
    phase2 = [
        ('EODHD integration', 'Replace yfinance with EODHD API for prices + fundamentals', 'High'),
        ('GICS industry groups', 'Upgrade from 11 sectors to ~25 industry groups', 'High'),
        ('Analyst estimates', 'Add EPFWD and EGRLF descriptors', 'High'),
        ('Scheduled updates', 'Daily data pipeline with Windows Task Scheduler', 'Medium'),
        ('Bias statistics', 'Implement CNE5 Appendix C backtesting framework', 'Medium'),
    ]
    add_table(doc, ['Task', 'Description', 'Priority'], phase2, col_widths=[4, 10, 3])

    doc.add_heading('9.3 Phase 3: Product Polish (Weeks 9-12)', level=2)
    phase3 = [
        ('Full Streamlit app', 'All 6 pages (upload, analysis, optimization, sizing, risk mgmt, status)', 'High'),
        ('Short-term reversal', 'Add reversal factor to style factors', 'Medium'),
        ('Specific risk model', 'Bayesian shrinkage for idiosyncratic risk', 'Medium'),
        ('Deploy to cloud', 'Railway or similar, with persistent storage', 'High'),
        ('User documentation', 'How to interpret outputs, FAQ, methodology guide', 'Medium'),
    ]
    add_table(doc, ['Task', 'Description', 'Priority'], phase3, col_widths=[4, 10, 3])

    doc.add_heading('9.4 Phase 4: Expansion (Weeks 13+)', level=2)
    phase4 = [
        ('European equities', 'STOXX 600 coverage as separate regional model', 'Medium'),
        ('Dash migration', 'If Streamlit limitations become blocking, migrate to Dash', 'Low'),
        ('Transaction costs', 'Spread + impact model for realistic optimization', 'Medium'),
        ('Portfolio tracking', 'Rebalancing framework, historical performance attribution', 'Medium'),
        ('Asia-Pacific equities', 'ASX 200, Nikkei 225, etc.', 'Low'),
    ]
    add_table(doc, ['Task', 'Description', 'Priority'], phase4, col_widths=[4, 10, 3])

    doc.add_page_break()

    # ================================================================
    # SECTION 10: GitHub Collaboration Guide
    # ================================================================
    doc.add_heading('10. GitHub Collaboration Guide', level=1)

    doc.add_paragraph(
        'This section explains how to set up efficient collaboration on GitHub '
        'for a two-person development team.'
    )

    doc.add_heading('10.1 Repository Setup', level=2)
    setup_steps = [
        '1. Create a GitHub repository (public or private) under one person\'s account.',
        '2. Add the collaborator: Settings \u2192 Collaborators \u2192 Add by username/email.',
        '3. Both clone the repo locally: git clone https://github.com/username/portfolio-xray.git',
        '4. Protect the main branch: Settings \u2192 Branches \u2192 Add rule for "main" \u2192 Require pull request reviews.',
    ]
    for step in setup_steps:
        doc.add_paragraph(step)

    doc.add_heading('10.2 Branch Strategy (GitHub Flow)', level=2)
    doc.add_paragraph(
        'Use the simple "GitHub Flow" model, which works well for small teams:'
    )
    branch_rules = [
        ('main', 'Always deployable. Never push directly. All changes via pull requests.'),
        ('feature/<name>', 'e.g., feature/streamlit-mvp, feature/eodhd-integration. One branch per task.'),
        ('fix/<name>', 'e.g., fix/beta-calculation. For bug fixes.'),
    ]
    add_table(doc, ['Branch', 'Purpose'], branch_rules, col_widths=[4, 13])

    doc.add_heading('10.3 Daily Workflow', level=2)
    workflow = [
        ('1. Pull latest', 'git checkout main && git pull'),
        ('2. Create branch', 'git checkout -b feature/my-task'),
        ('3. Work & commit', 'git add <files> && git commit -m "Add description"'),
        ('4. Push branch', 'git push -u origin feature/my-task'),
        ('5. Open PR', 'On GitHub: click "Compare & pull request", describe changes'),
        ('6. Code review', 'Partner reviews, comments, approves'),
        ('7. Merge', 'Click "Squash and merge" on GitHub'),
        ('8. Clean up', 'git checkout main && git pull && git branch -d feature/my-task'),
    ]
    add_table(doc, ['Step', 'Command / Action'], workflow, col_widths=[3, 14])

    doc.add_heading('10.4 Recommended Tools', level=2)
    tools = [
        ('GitHub Issues', 'Track tasks, bugs, and feature requests. Use labels (bug, enhancement, data, frontend). Free.'),
        ('GitHub Projects', 'Kanban board built into GitHub. Columns: Backlog, In Progress, Review, Done. Free.'),
        ('GitHub Actions', 'CI/CD: automatically run tests on every PR. Free for public repos, 2000 min/month for private.'),
        ('Claude Code', 'AI coding assistant for implementation, debugging, and code review. Both developers can use it.'),
        ('VS Code + Live Share', 'Real-time collaborative editing for pair programming sessions.'),
    ]
    add_table(doc, ['Tool', 'Use Case'], tools, col_widths=[4, 13])

    doc.add_heading('10.5 Communication', level=2)
    doc.add_paragraph(
        'For a 2-person team, keep it simple:'
    )
    comm = [
        'Use GitHub Issues as the single source of truth for what needs to be done',
        'Use PR comments for code-specific discussions (keeps context with the code)',
        'Use WhatsApp/Discord/Slack for quick questions and coordination',
        'Weekly 30-min sync to review progress and plan next week\'s tasks',
    ]
    for item in comm:
        doc.add_paragraph(item, style='List Bullet')

    doc.add_heading('10.6 Dividing Work', level=2)
    doc.add_paragraph(
        'To minimize merge conflicts, divide work by component:'
    )
    division = [
        ('Person A', 'Backend: factor model, data pipeline, EODHD integration, tests'),
        ('Person B', 'Frontend: Streamlit pages, DuckDB integration, deployment, documentation'),
        ('Shared', 'Architecture decisions, code reviews, Excel verification model'),
    ]
    add_table(doc, ['Who', 'Responsibility'], division, col_widths=[3, 14])

    # Save
    output_path = 'Portfolio_XRay_Project_Plan.docx'
    doc.save(output_path)
    print(f'Saved: {output_path}')
    return output_path


if __name__ == '__main__':
    build_document()
